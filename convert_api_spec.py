import sys
import os
import subprocess

try:
    import docx
except ImportError:
    subprocess.check_call([sys.executable, "-m", "pip", "install", "python-docx"])
    import docx

from docx import Document
from docx.shared import Pt, Inches

md_path = r'c:\Users\ubee2\Downloads\DAE_P1_19Modules-feature-refactor-simulation-m17\DAE_P1_19Modules-feature-refactor-simulation-m17\DOCS\API_SPEC_ZH.md'
out_path = r'c:\Users\ubee2\Downloads\DAE_P1_19Modules-feature-refactor-simulation-m17\DAE_P1_19Modules-feature-refactor-simulation-m17\DOCS\API_SPEC_ZH.docx'

with open(md_path, 'r', encoding='utf-8') as f:
    lines = f.readlines()

doc = Document()

in_code_block = False

for line in lines:
    line = line.rstrip('\n')
    
    if line.startswith('```'):
        in_code_block = not in_code_block
        continue
        
    if in_code_block:
        p = doc.add_paragraph(line)
        p.paragraph_format.left_indent = Inches(0.5)
        for run in p.runs:
            run.font.name = 'Courier New'
            run.font.size = Pt(9)
        continue

    if line.startswith('# '):
        doc.add_heading(line[2:], 0)
    elif line.startswith('## '):
        doc.add_heading(line[3:], level=1)
    elif line.startswith('### '):
        doc.add_heading(line[4:], level=2)
    elif line.startswith('- **'):
        p = doc.add_paragraph(style='List Bullet')
        parts = line[4:].split('**', 1)
        if len(parts) == 2:
            r1 = p.add_run(parts[0])
            r1.bold = True
            p.add_run(parts[1])
        else:
            p.add_run(line[2:])
    elif line.startswith('- '):
        doc.add_paragraph(line[2:], style='List Bullet')
    elif line.strip() == '---':
        continue
    elif line.strip():
        doc.add_paragraph(line)

doc.save(out_path)
print(f"Successfully generated {out_path}")
