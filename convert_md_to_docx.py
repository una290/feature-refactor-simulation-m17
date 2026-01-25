import re
import sys
from docx import Document
from docx.shared import Pt
import os

def parse_markdown_table(lines):
    """
    Parse a list of markdown table lines into a list of lists (rows/cols).
    Skips the separator line (e.g. |---|---|).
    """
    table_data = []
    for line in lines:
        if set(line.strip()) <= {'|', '-', ' ', ':'}:
            continue
        # Split by pipe, remove first/last empty strings if they exist
        cells = [c.strip() for c in line.strip().split('|')]
        # Markdown tables usually start/end with |, so split gives empty first/last
        if cells and cells[0] == '':
            cells.pop(0)
        if cells and cells[-1] == '':
            cells.pop()
        table_data.append(cells)
    return table_data

def add_markdown_paragraph(doc, text, style='Body Text'):
    """
    Add a paragraph with basic inline bold formatting support.
    """
    # Simply mapping styles
    if style.startswith('Heading'):
        p = doc.add_heading(level=int(style[-1]))
    else:
        p = doc.add_paragraph(style='List Bullet' if style == 'List Bullet' else None)

    # Split by bold markers
    # **text**
    parts = re.split(r'(\*\*[^*]+\*\*)', text)
    for part in parts:
        if part.startswith('**') and part.endswith('**'):
            run = p.add_run(part[2:-2])
            run.bold = True
        else:
            p.add_run(part)

def convert_md_to_docx(md_path, docx_path):
    print(f"Reading {md_path}...")
    with open(md_path, 'r', encoding='utf-8') as f:
        lines = f.readlines()

    doc = Document()
    
    # Simple line iterator
    i = 0
    while i < len(lines):
        line = lines[i].rstrip()
        
        # Skip empty lines
        if not line:
            i += 1
            continue

        # Headers
        if line.startswith('#'):
            level = len(line.split(' ')[0])
            text = line[level:].strip()
            add_markdown_paragraph(doc, text, style=f'Heading {level}')
            i += 1
            continue
            
        # List items
        if line.lstrip().startswith('- '):
            text = line.lstrip()[2:].strip()
            add_markdown_paragraph(doc, text, style='List Bullet')
            i += 1
            continue
            
        # Tables
        if line.strip().startswith('|'):
            # Collect all table lines
            table_lines = []
            while i < len(lines) and lines[i].strip().startswith('|'):
                table_lines.append(lines[i])
                i += 1
            
            data = parse_markdown_table(table_lines)
            if data:
                rows = len(data)
                cols = len(data[0])
                table = doc.add_table(rows=rows, cols=cols)
                table.style = 'Table Grid'
                for r_idx, row_data in enumerate(data):
                    row = table.rows[r_idx]
                    for c_idx, cell_text in enumerate(row_data):
                        # Use same logic to bold inside table cells
                        if c_idx < len(row.cells):
                            # Clear default par
                            row.cells[c_idx].text = ""
                            p = row.cells[c_idx].paragraphs[0]
                            parts = re.split(r'(\*\*[^*]+\*\*)', cell_text)
                            for part in parts:
                                if part.startswith('**') and part.endswith('**'):
                                    run = p.add_run(part[2:-2])
                                    run.bold = True
                                else:
                                    p.add_run(part)
            continue

        # Default Paragraph
        text = line.strip()
        add_markdown_paragraph(doc, text, style='Normal')
        i += 1

    print(f"Saving to {docx_path}...")
    doc.save(docx_path)
    print("Done.")

if __name__ == "__main__":
    md_file = "DOCS/custom_parameters_zh_TW.md"
    docx_file = "DOCS/custom_parameters_zh_TW.docx"
    
    # Adjust paths to absolute if needed, assuming running from project root
    base_dir = r"C:\Users\ubee2\Downloads\DAE_P1_19Modules-feature-refactor-simulation-m17\DAE_P1_19Modules-feature-refactor-simulation-m17"
    
    md_full = os.path.join(base_dir, md_file)
    docx_full = os.path.join(base_dir, docx_file)
    
    convert_md_to_docx(md_full, docx_full)
