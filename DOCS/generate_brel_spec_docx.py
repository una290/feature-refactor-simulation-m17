"""
將 BREL_Spec_v2.md 轉換為 BREL_Spec_v2.docx
使用 python-docx 進行格式化輸出
"""
import os
import re
import sys
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

def add_heading(doc, text, level=1):
    h = doc.add_heading(text, level=level)
    h.runs[0].font.color.rgb = RGBColor(0x1a, 0x23, 0x7e) if level == 1 else RGBColor(0x01, 0x57, 0x9b)
    return h

def add_code_block(doc, code_text):
    """新增程式碼區塊（灰底、等寬字體）"""
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.3)
    pPr = p._p.get_or_add_pPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), 'F5F5F5')
    pPr.append(shd)
    run = p.add_run(code_text)
    run.font.name = 'Courier New'
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(0x2e, 0x7d, 0x32)

def add_table_from_md(doc, lines):
    """從 Markdown table 行解析並建立 Word 表格"""
    rows_data = []
    for line in lines:
        if line.strip().startswith('|') and not re.match(r'^\|[-:| ]+\|$', line.strip()):
            cells = [c.strip() for c in line.strip().strip('|').split('|')]
            rows_data.append(cells)
    if not rows_data:
        return
    col_count = max(len(r) for r in rows_data)
    table = doc.add_table(rows=0, cols=col_count)
    table.style = 'Table Grid'
    for i, row_data in enumerate(rows_data):
        row = table.add_row()
        for j, cell_text in enumerate(row_data):
            if j < col_count:
                cell = row.cells[j]
                cell.text = cell_text
                if i == 0:
                    for run in cell.paragraphs[0].runs:
                        run.font.bold = True
                        run.font.color.rgb = RGBColor(0xff, 0xff, 0xff)
                    tc = cell._tc
                    tcPr = tc.get_or_add_tcPr()
                    shd = OxmlElement('w:shd')
                    shd.set(qn('w:val'), 'clear')
                    shd.set(qn('w:color'), 'auto')
                    shd.set(qn('w:fill'), '1A237E')
                    tcPr.append(shd)
    doc.add_paragraph()

def md_to_docx(md_text, doc):
    lines = md_text.splitlines()
    i = 0
    table_buffer = []
    code_buffer = []
    in_code = False
    in_table = False

    while i < len(lines):
        line = lines[i]

        # 程式碼區塊
        if line.strip().startswith('```'):
            if not in_code:
                in_code = True
                code_buffer = []
            else:
                in_code = False
                add_code_block(doc, '\n'.join(code_buffer))
            i += 1
            continue

        if in_code:
            code_buffer.append(line)
            i += 1
            continue

        # 表格偵測
        if line.strip().startswith('|'):
            if not in_table:
                in_table = True
                table_buffer = []
            table_buffer.append(line)
            i += 1
            continue
        else:
            if in_table:
                add_table_from_md(doc, table_buffer)
                in_table = False
                table_buffer = []

        # 標題
        if line.startswith('# '):
            add_heading(doc, line[2:].strip(), level=1)
        elif line.startswith('## '):
            add_heading(doc, line[3:].strip(), level=2)
        elif line.startswith('### '):
            add_heading(doc, line[4:].strip(), level=3)
        elif line.startswith('#### '):
            add_heading(doc, line[5:].strip(), level=4)
        # 引用塊
        elif line.startswith('> [!'):
            tag_match = re.match(r'> \[!(\w+)\]', line)
            tag = tag_match.group(1) if tag_match else 'NOTE'
            color_map = {
                'NOTE': RGBColor(0x01, 0x57, 0x9b),
                'WARNING': RGBColor(0xe6, 0x51, 0x00),
                'IMPORTANT': RGBColor(0x1b, 0x5e, 0x20),
                'CAUTION': RGBColor(0xb7, 0x1c, 0x1c),
                'TIP': RGBColor(0x33, 0x69, 0x1e),
            }
            color = color_map.get(tag, RGBColor(0x55, 0x55, 0x55))
            note_lines = []
            i += 1
            while i < len(lines) and lines[i].startswith('>'):
                note_lines.append(lines[i].lstrip('> ').strip())
                i += 1
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.3)
            run = p.add_run(f"[{tag}] " + ' '.join(note_lines))
            run.font.color.rgb = color
            run.font.italic = True
            continue
        # 分隔線
        elif line.strip() == '---':
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(6)
        # 清單
        elif re.match(r'^(\s*)[-*] ', line):
            p = doc.add_paragraph(style='List Bullet')
            text = re.sub(r'^(\s*)[-*] ', '', line)
            text = re.sub(r'\*\*(.+?)\*\*', r'\1', text)
            p.add_run(text)
        # 一般段落
        elif line.strip():
            clean = re.sub(r'\*\*(.+?)\*\*', r'\1', line)
            clean = re.sub(r'`(.+?)`', r'\1', clean)
            doc.add_paragraph(clean.strip())

        i += 1

    if in_table and table_buffer:
        add_table_from_md(doc, table_buffer)

def main():
    SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
    INPUT_MD = os.path.join(SCRIPT_DIR, "BREL_Spec_v2.md")
    OUTPUT_DOCX = os.path.join(SCRIPT_DIR, "BREL_Spec_v2.docx")

    if not os.path.exists(INPUT_MD):
        print(f"Error: {INPUT_MD} not found.")
        return

    with open(INPUT_MD, "r", encoding="utf-8") as f:
        md_text = f.read()

    doc = Document()
    
    # 頁面設定
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)

    md_to_docx(md_text, doc)
    
    try:
        doc.save(OUTPUT_DOCX)
        print(f"Successfully saved to {OUTPUT_DOCX}")
    except PermissionError:
        ALT_OUTPUT = os.path.join(SCRIPT_DIR, "BREL_Spec_v2_updated.docx")
        doc.save(ALT_OUTPUT)
        print(f"Permission denied for {OUTPUT_DOCX}. Saved to {ALT_OUTPUT} instead.")
    except Exception as e:
        print(f"Error saving DOCX: {e}")

if __name__ == "__main__":
    main()
